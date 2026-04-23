"""감사 로그 조회 — ADMIN / VIEWER (감사관 역할 겸함) 만 접근."""

from __future__ import annotations

import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import and_, select
from sqlalchemy.orm import Session

from app.api.deps import get_db, require_role
from app.models import AuditLog, UserRole
from app.schemas.audit import AuditLogOut

router = APIRouter(prefix="/audit-logs", tags=["audit-logs"])

# 감사 로그는 운영자(OPERATOR) 도 못 봄. ADMIN + VIEWER (감사관 역할) 만.
_AUDITOR_ROLES = [UserRole.ADMIN, UserRole.VIEWER]


def _filters(
    actor: str | None,
    action: str | None,
    resource_type: str | None,
    since: datetime | None,
    until: datetime | None,
):
    conds = []
    if actor:
        conds.append(AuditLog.actor_username.ilike(f"%{actor}%"))
    if action:
        conds.append(AuditLog.action.ilike(f"%{action}%"))
    if resource_type:
        conds.append(AuditLog.resource_type == resource_type)
    if since is not None:
        conds.append(AuditLog.occurred_at >= since)
    if until is not None:
        conds.append(AuditLog.occurred_at < until)
    return and_(*conds) if conds else None


@router.get("", response_model=list[AuditLogOut])
def list_audit_logs(
    actor: str | None = None,
    action: str | None = None,
    resource_type: str | None = None,
    since: datetime | None = None,
    until: datetime | None = None,
    limit: int = Query(200, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    _=Depends(require_role(_AUDITOR_ROLES)),
) -> list[AuditLogOut]:
    stmt = select(AuditLog)
    where = _filters(actor, action, resource_type, since, until)
    if where is not None:
        stmt = stmt.where(where)
    stmt = stmt.order_by(AuditLog.occurred_at.desc()).limit(limit).offset(offset)
    return [AuditLogOut.model_validate(r) for r in db.scalars(stmt).all()]


@router.get("/actions", response_model=list[str])
def list_actions(
    db: Session = Depends(get_db),
    _=Depends(require_role(_AUDITOR_ROLES)),
) -> list[str]:
    """필터 셀렉트용 — DB에 실제로 존재하는 action 코드만 반환."""
    return [
        r[0]
        for r in db.execute(select(AuditLog.action).distinct().order_by(AuditLog.action))
    ]


@router.get("/resource-types", response_model=list[str])
def list_resource_types(
    db: Session = Depends(get_db),
    _=Depends(require_role(_AUDITOR_ROLES)),
) -> list[str]:
    return [
        r[0]
        for r in db.execute(
            select(AuditLog.resource_type)
            .where(AuditLog.resource_type.is_not(None))
            .distinct()
            .order_by(AuditLog.resource_type)
        )
    ]


@router.get("/export.xlsx")
def export_xlsx(
    actor: str | None = None,
    action: str | None = None,
    resource_type: str | None = None,
    since: datetime | None = None,
    until: datetime | None = None,
    db: Session = Depends(get_db),
    _=Depends(require_role(_AUDITOR_ROLES)),
) -> StreamingResponse:
    """현재 필터 조건의 감사 로그를 Excel 한 시트에 덤프."""
    import json as _json

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    stmt = select(AuditLog)
    where = _filters(actor, action, resource_type, since, until)
    if where is not None:
        stmt = stmt.where(where)
    rows = db.scalars(stmt.order_by(AuditLog.occurred_at.desc()).limit(10000)).all()

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "감사 로그"
    ws.append([
        "ID", "시각(KST)", "사용자", "역할", "액션",
        "자원 종류", "자원 ID", "변경 전", "변경 후", "IP", "User-Agent",
    ])
    header_fill = PatternFill("solid", fgColor="1F3A93")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r.id,
            r.occurred_at.strftime("%Y-%m-%d %H:%M:%S") if r.occurred_at else "",
            r.actor_username,
            r.actor_role or "",
            r.action,
            r.resource_type or "",
            r.resource_id or "",
            _json.dumps(r.before_value, ensure_ascii=False) if r.before_value else "",
            _json.dumps(r.after_value, ensure_ascii=False) if r.after_value else "",
            r.ip or "",
            r.user_agent or "",
        ])
    for col_letter, width in zip("ABCDEFGHIJK", [8, 20, 14, 12, 26, 14, 12, 50, 50, 16, 30]):
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"audit-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )