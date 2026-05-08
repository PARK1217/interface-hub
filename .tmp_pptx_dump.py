from pptx import Presentation
prs = Presentation('NOA_Interface_Hub_기획서 (1).pptx')
for i, slide in enumerate(prs.slides, 1):
    print(f'=== Slide {i} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = ''.join(r.text for r in para.runs).strip()
                if t:
                    print(f'  {t}')
        if shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip().replace(chr(10), ' / ') for cell in row.cells]
                print('  [TABLE]', ' | '.join(cells))
    print()
